"""Conversión del Informe General (Excel de la empresa) al formato MATRIZ.

1. Detecta las filas pintadas de rojo/rosado y las elimina.
2. Conserva las demás filas.
3. Genera un Excel con las columnas del modelo MATRIZ y el aspecto del informe
   (título, fecha del documento, encabezado verde).
"""
from __future__ import annotations

import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from cleaner import DamagedFileError

# ------------------------------------------------------------------ modelo
# Columnas del modelo MATRIZ (hoja MATRIZ del Excel MATRIZ_RESIDUOS_2026)
MATRIZ_HEADERS = [
    "FECHA SERVICIO", "MANIFIESTO SUMMA", "ACTA TEMPORAL SUMMA", "CLIENTE", "DIRECCIÓN",
    "RESIDUO", "TRATAMIENTO", "PESO", "UNIDAD DE MEDIDA", "PROVEEDOR", "ESTADO",
    "PESO PENDIENTE", "PESO TIQUETE", "PESO REAL PENDIENTE", "FECHA DISPOSICIÓN FINAL",
    "MES DE DISPOSICIÓN", "REGISTRO DE INGRESO DISPOSITOR FINAL (RM, TRA, DRS)", "CONCILIADO",
    "CERTIFICADO", "ACTASFINALES", "CARGADO EN PLATAFORMA AL CLIENTE", "COMENTARIOS",
]
TITLE = "MATRIZ GESTIÓN RESIDUOS CLIENTES - SOLUCIÓN EMPRESARIAL SUMMA SAS"

GREEN_DARK = "1D7E4D"
GREEN_TEXT = "62A75F"

# Encabezados del informe que se necesitan (en minúsculas)
NEEDED = {
    "fecha": "fecha servicio",
    "manifiesto": "manifiesto/remisión",
    "acta": "acta",
    "cliente": "cliente/gestor",
    "direccion": "dirección sucursal",
    "residuo": "residuo/servicio",
    "tratamiento": "tipo de disposición",
    "peso_conc": "cantidad conciliado",
    "peso_carg": "cantidad cargado",
    "unidad": "unidad conciliado",
    "unidad_carg": "unidad programado/cargado",
    "proveedor": "prov/dest",
}

TRATAMIENTOS = {
    "RELLENO EN ESCOMBRERA": "ESCOMBRERA",
    "OSMOSIS INVERSA": "OSMOSIS CON ASPERSIÓN AL SUELO",
    "INCINERACION": "INCINERACIÓN",
    "LAMINAS FILTRANTES": "LAMINAS FILTRANTES",
}


# --------------------------------------------------------------- utilidades
def is_red_fill(cell) -> bool:
    """True si la celda tiene relleno rojo o rosado (incluye el rosa claro tipo FDB5B6)."""
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return False
    rgb = fill.fgColor.rgb if fill.fgColor.type == "rgb" else None
    if not isinstance(rgb, str) or len(rgb) < 6:
        return False
    r, g, b = (int(rgb[-6:][i : i + 2], 16) for i in (0, 2, 4))
    return r >= 200 and (r - g) >= 40 and (r - b) >= 40


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _tratamiento(v):
    v = _clean(v)
    if not isinstance(v, str):
        return v
    base = v.split("|")[0].strip().upper()
    return TRATAMIENTOS.get(base, base)


def _proveedor(v):
    v = _clean(v)
    if not isinstance(v, str):
        return v
    up = v.upper()
    if "BIOAGRICOLA" in up:
        return "BIOAGRICOLA"
    if "BIO PROYECTOS" in up or "BIOPROYECTOS" in up:
        return "BIOPROYECTOS"
    if "SUMMA" in up:
        return "SOLUCIÓN EMPRESARIAL SUMMA"
    return up.split("-SEDE")[0].strip()


def _to_date(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(v.strip()[:10], fmt)
            except ValueError:
                pass
    return None


# ------------------------------------------------------------------ lectura
def read_informe(file_bytes: bytes) -> dict:
    """Lee el informe y separa filas rojas (a borrar) de filas a conservar."""
    if not file_bytes:
        raise DamagedFileError("El archivo está completamente vacío (0 bytes).")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
        ws = wb.active
    except Exception as exc:  # noqa: BLE001
        raise DamagedFileError(f"No se pudo leer el archivo: {exc}") from exc

    # fila de encabezados: la que contiene "Fecha servicio"
    header_row = None
    for r in range(1, min(ws.max_row, 20) + 1):
        if any(str(c.value or "").strip().lower() == NEEDED["fecha"] for c in ws[r]):
            header_row = r
            break
    if header_row is None:
        raise DamagedFileError("No se encontró la fila de encabezados ('Fecha servicio'). ¿Es el Informe General?")

    cols = {}
    for c in ws[header_row]:
        name = str(c.value or "").strip().lower()
        for key, wanted in NEEDED.items():
            if name == wanted and key not in cols:
                cols[key] = c.column
    missing = [NEEDED[k] for k in ("fecha", "manifiesto", "cliente", "residuo") if k not in cols]
    if missing:
        raise DamagedFileError(f"Faltan columnas necesarias en el informe: {', '.join(missing)}")

    keep, deleted = [], []
    for r in range(header_row + 1, ws.max_row + 1):
        fecha = _to_date(ws.cell(r, cols["fecha"]).value)
        if fecha is None:  # fila de totales, vacía, etc.
            continue
        get = lambda k: ws.cell(r, cols[k]).value if k in cols else None  # noqa: E731
        peso = get("peso_conc")
        unidad = get("unidad")
        if peso in (None, ""):
            peso, unidad = get("peso_carg"), get("unidad_carg")
        record = {
            "excel_fila": r,
            "fecha": fecha,
            "manifiesto": get("manifiesto"),
            "acta": _clean(get("acta")),
            "cliente": _clean(get("cliente")),
            "direccion": _clean(get("direccion")),
            "residuo": _clean(get("residuo")),
            "tratamiento": _tratamiento(get("tratamiento")),
            "peso": peso if isinstance(peso, (int, float)) else None,
            "unidad": _clean(unidad),
            "proveedor": _proveedor(get("proveedor")),
        }
        # la fila es "roja" si la celda de Fecha servicio está pintada de rojo/rosado
        (deleted if is_red_fill(ws.cell(r, cols["fecha"])) else keep).append(record)

    return {"keep": keep, "deleted": deleted, "total": len(keep) + len(deleted)}


# ---------------------------------------------------------------- escritura
def build_matriz(records: list[dict], deleted: list[dict] | None = None) -> bytes:
    """Excel con columnas del modelo MATRIZ y el aspecto del informe (foto)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MATRIZ"
    n_cols = len(MATRIZ_HEADERS)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    ws["A1"] = TITLE
    ws["A1"].font = Font(name="Verdana", size=14, bold=True, color=GREEN_DARK)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    ws["A2"] = f"Fecha documento: {datetime.now():%Y-%m-%d %H:%M:%S}"
    ws["A2"].font = Font(name="Verdana", size=10, bold=True, color=GREEN_TEXT)

    head_fill = PatternFill("solid", fgColor=GREEN_DARK)
    for i, h in enumerate(MATRIZ_HEADERS, start=1):
        c = ws.cell(3, i, h)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = head_fill
        c.border = border
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 32

    base = Font(name="Calibri", size=11)
    for n, rec in enumerate(records):
        r = 4 + n
        values = {
            1: rec["fecha"], 2: rec["manifiesto"], 3: rec["acta"], 4: rec["cliente"],
            5: rec["direccion"], 6: rec["residuo"], 7: rec["tratamiento"], 8: rec["peso"],
            9: rec["unidad"], 10: rec["proveedor"],
            # fórmulas del modelo MATRIZ
            11: f'=IF(O{r}="","PENDIENTE","DISPUESTO")',
            12: f'=IF(K{r}="DISPUESTO",0,H{r})',
            14: f"=H{r}-M{r}",
            16: f'=IF(ISNUMBER(O{r}),TEXT(O{r},"MMMM"),"")',
        }
        for col in range(1, n_cols + 1):
            c = ws.cell(r, col, values.get(col))
            c.font = base
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=col in (4, 5, 6, 7))
        ws.cell(r, 1).number_format = "yyyy/mm/dd"
        ws.cell(r, 1).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(r, 2).alignment = Alignment(horizontal="right", vertical="center")
        for col in (8, 12, 13, 14):
            ws.cell(r, col).number_format = "#,##0.00"
        ws.cell(r, 15).number_format = "yyyy/mm/dd"

    widths = [16, 18, 20, 38, 42, 38, 34, 12, 12, 28, 14, 16, 14, 18, 22, 18, 30, 14, 14, 22, 20, 28]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(n_cols)}{3 + max(len(records), 1)}"

    if deleted:
        wd = wb.create_sheet("Eliminadas (rojas)")
        heads = ["Fila en el informe", "Fecha servicio", "Manifiesto", "Cliente", "Residuo"]
        for i, h in enumerate(heads, start=1):
            c = wd.cell(1, i, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
        for n, rec in enumerate(deleted, start=2):
            wd.cell(n, 1, rec["excel_fila"])
            wd.cell(n, 2, rec["fecha"]).number_format = "yyyy/mm/dd"
            wd.cell(n, 3, rec["manifiesto"])
            wd.cell(n, 4, rec["cliente"])
            wd.cell(n, 5, rec["residuo"])
        for i, w in enumerate([18, 16, 14, 40, 40], start=1):
            wd.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
